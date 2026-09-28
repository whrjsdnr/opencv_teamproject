"""Build final Markdown/PPTX/PDF from frozen measured CSV; no hand-entered metrics."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path
import shutil
import textwrap


from scripts import run_experiments as exp

ROOT = exp.ROOT
OUT = ROOT / "outputs/final_report"
FIG = OUT / "figures"
TAB = OUT / "tables"


def numeric_rows(path):
    rows = exp.read_csv(path)
    for row in rows:
        for key, value in row.items():
            parsed = exp.number(value)
            if math.isfinite(parsed):
                row[key] = parsed
    return rows


def aggregate(rows):
    tp, fp, fn = (sum(r[k] for r in rows) for k in ("tp", "fp", "fn"))
    count = sum(r["alert_delay_count"] for r in rows)
    return dict(
        tp=tp,
        fp=fp,
        fn=fn,
        precision=tp / (tp + fp) if tp + fp else exp.NAN,
        recall=tp / (tp + fn) if tp + fn else exp.NAN,
        f1=2 * tp / (2 * tp + fp + fn) if 2 * tp + fp + fn else exp.NAN,
        fps=exp.mean(r["fps"] for r in rows),
        alert_delay=sum(r["alert_delay_sum"] for r in rows) / count if count else exp.NAN,
    )


def test_count():
    text = (TAB / "pytest.txt").read_text()
    match = re.search(r"(\d+) passed", text)
    if not match:
        raise ValueError("A passing pytest result is required")
    return int(match.group(1))


def fmt(value, digits=3):
    if isinstance(value, (int, float)):
        return f"{value:.{digits}f}" if math.isfinite(value) else "N/A"
    return str(value)


def figure(data, key, labels, name, ylabel):
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update({"font.size": 11, "axes.spines.top": False, "axes.spines.right": False})
    fig, ax = plt.subplots(figsize=(9, 4.5))
    vals = [r[key] for r in data]
    bars = ax.bar(labels, vals, color=["#20A69A", "#3479B9", "#F0AB4B"] * (len(vals) // 3 + 1))
    for bar, value in zip(bars, vals):
        if math.isfinite(value):
            ax.annotate(
                fmt(value, 2),
                (bar.get_x() + bar.get_width() / 2, value),
                ha="center",
                va="bottom",
                fontsize=9,
            )
    ax.set_ylabel(ylabel)
    ax.tick_params(axis="x", labelrotation=20 if len(labels) > 3 else 0)
    fig.tight_layout()
    fig.savefig(FIG / name, dpi=180)
    plt.close(fig)


def representative_frame(row):
    """Replay real backend to export its actual UI; never use this replay for reported FPS."""
    import cv2
    from src.pipeline import RealBackend, _draw, validate_config

    environment = json.loads((Path(row["run_dir"]) / "effective_config.json").read_text())
    config = validate_config(environment["config"])
    video = ROOT / "data/samples" / row["video"]
    cap = cv2.VideoCapture(str(video))
    native = tuple(environment["native_size"])
    backend = RealBackend(config, native, native)
    fps = cap.get(cv2.CAP_PROP_FPS)
    index = 0
    selected = None
    try:
        while True:
            ok, frame = cap.read()
            if not ok:
                break
            packet = dict(frame_index=index, media_time_s=index / fps)
            intrusions, boxes, _ = backend.analyze(frame, packet)
            states, transitions = backend.update(intrusions, packet)
            if index == 0:
                cv2.imwrite(
                    str(FIG / "zone_view.png"),
                    _draw(frame, "baseline", 0, backend.zones, boxes, states, False, False),
                )
            if index / fps >= 2 and any(s["status"] == "ALERT" for s in states.values()):
                cv2.imwrite(
                    str(FIG / "actual_alert.png"),
                    _draw(frame, "baseline", 0, backend.zones, boxes, states, True, False),
                )
                selected = dict(
                    video=row["video"],
                    frame_index=index,
                    media_time_s=index / fps,
                    provenance="RealBackend replay using frozen effective config; display FPS=0 (not a measurement)",
                    run_dir=row["run_dir"],
                )
                break
            index += 1
    finally:
        cap.release()
    (TAB / "screenshot_provenance.json").write_text(json.dumps(selected, indent=2))


def prepare(batch, study):
    rows = numeric_rows(batch / "summary/experiment_summary.csv")
    modes = numeric_rows(batch / "summary/mode_summary.csv")
    studies = numeric_rows(study / "study_results.csv")
    assert len(rows) == 63, "Expected 7 videos x 3 modes x 3 repeats"
    assert len(studies) == 76, "Expected 19 cases x 2 videos x 2 repeats"
    assert all(
        r["status"] == "COMPLETED" and r["evaluation_status"] == "EVALUATED" for r in rows + studies
    )
    assert all(
        r["recording_status"] == "PASS" and r["pending_at_stop"] == 0 for r in rows + studies
    )
    for r in rows:
        assert r["processed_frames"] == r["frame_count"] and r["dropped_frames"] == 0
    before = json.loads((TAB / "input_integrity_before.json").read_text())
    assert all(
        hashlib.sha256(Path(p).read_bytes()).hexdigest() == digest for p, digest in before.items()
    )
    for path in (batch / "summary").glob("*.csv"):
        shutil.copy2(path, TAB / path.name)
    shutil.copy2(study / "study_results.csv", TAB / "study_results.csv")
    for path in (batch / "summary/figures").glob("*.png"):
        shutil.copy2(path, FIG / path.name)
    videos = []
    for video in sorted({r["video"] for r in rows}):
        for mode in exp.MODES:
            selected = [r for r in rows if r["video"] == video and r["mode"] == mode]
            summary = aggregate(selected)
            videos.append(
                dict(
                    video=video,
                    mode=mode,
                    repeats=len(selected),
                    **summary,
                    open_at_eof_count=sum(r["open_at_eof_count"] for r in selected),
                )
            )
    resolution = []
    for size in ("192x144", "288x216", "384x288"):
        for mode in exp.MODES:
            selected = [
                r
                for r in studies
                if r["study"] == "resolution" and r["setting"] == size and r["mode"] == mode
            ]
            resolution.append(
                dict(resolution=size, mode=mode, runs=len(selected), **aggregate(selected))
            )
    optimization = []
    for technique in dict.fromkeys(r["technique"] for r in studies if r["study"] == "optimization"):
        off, on = [
            aggregate(
                [r for r in studies if r["technique"] == technique and r["setting"] == setting]
            )
            for setting in ("OFF", "ON")
        ]
        optimization.append(
            dict(
                technique=technique,
                off_fps=off["fps"],
                on_fps=on["fps"],
                improvement_percent=exp.improvement(on["fps"], off["fps"]),
                off_f1=off["f1"],
                on_f1=on["f1"],
                delta_f1=on["f1"] - off["f1"],
                off_tp=off["tp"],
                off_fp=off["fp"],
                off_fn=off["fn"],
                on_tp=on["tp"],
                on_fp=on["fp"],
                on_fn=on["fn"],
            )
        )
    stage_data = numeric_rows(batch / "summary/stage_timings.csv")
    stages = [
        dict(
            mode=mode,
            stage=stage,
            mean_ms=exp.mean(
                r["mean_ms"] for r in stage_data if r["mode"] == mode and r["stage"] == stage
            ),
        )
        for mode in exp.MODES
        for stage in exp.STAGES
    ]
    for name, values in [
        ("video_metrics", videos),
        ("resolution_comparison", resolution),
        ("optimization_comparison", optimization),
        ("stage_aggregate", stages),
    ]:
        exp.save_csv(TAB / f"{name}.csv", values)
    figure(
        resolution,
        "fps",
        [r["resolution"] + "\n" + r["mode"] for r in resolution],
        "resolution_matrix.png",
        "Processing throughput FPS (2 videos x 2 repeats)",
    )
    figure(
        optimization,
        "improvement_percent",
        [r["technique"] for r in optimization],
        "optimization_effects.png",
        "ON vs OFF throughput change (%)",
    )
    direct = [
        r
        for r in stages
        if r["mode"] == "baseline" and r["stage"] not in ("total", "queue_wait", "display")
    ]
    figure(
        direct,
        "mean_ms",
        [r["stage"] for r in direct],
        "bottleneck.png",
        "Baseline mean stage time (ms)",
    )
    representative_frame(
        next(r for r in rows if r["video"] == "Danger.mp4" and r["mode"] == "baseline")
    )
    latency = [
        dict(
            mode=mode,
            processed_frames=sum(r["processed_frames"] for r in rows if r["mode"] == mode),
            measured_frames=sum(r["measured_frames"] for r in rows if r["mode"] == mode),
            elapsed_seconds=sum(r["elapsed_time"] for r in rows if r["mode"] == mode),
            queue_wait_ms=exp.mean(r["queue_wait_ms"] for r in rows if r["mode"] == mode),
            total_frame_ms=exp.mean(r["total_frame_ms"] for r in rows if r["mode"] == mode),
            dropped_frames=sum(r["dropped_frames"] for r in rows if r["mode"] == mode),
        )
        for mode in exp.MODES
    ]
    exp.save_csv(TAB / "performance_totals.csv", latency)
    matches = numeric_rows(batch / "summary/event_matches.csv")
    failures = [
        m
        for m in matches
        if m["mode"] == "baseline" and m["repeat"] == 1 and m["outcome"] in ("FP", "FN")
    ]
    truth_rows = numeric_rows(batch / "summary/ground_truth_normalized.csv")
    truth_rows = [
        r
        for r in truth_rows
        if r["mode"] == "baseline" and r["repeat"] == 1 and r["event_type"] == "intrusion"
    ]
    for failure in failures:
        if failure["outcome"] == "FN":
            failure["analysis"] = "해당 GT에 매칭된 경보 없음"
        else:
            overlaps = [
                g
                for g in truth_rows
                if g["video_name"] == failure["video"]
                and g["zone_name"] == failure["zone"]
                and failure["detected_start"] <= g["end_time"]
                and failure["evaluation_end_time"] >= g["start_time"]
            ]
            failure["analysis"] = (
                "GT와 겹치지만 이미 1:1 배정된 중복 경보"
                if overlaps
                else "어떤 GT 구간과도 겹치지 않는 경보"
            )
    exp.save_csv(TAB / "failure_cases.csv", failures)
    return dict(
        batch=batch,
        study=study,
        failures=failures,
        rows=rows,
        modes=modes,
        videos=videos,
        resolution=resolution,
        optimization=optimization,
        stages=stages,
    )


def make_markdown(data):
    batch, study = data["batch"], data["study"]
    rows, modes, videos = data["rows"], data["modes"], data["videos"]
    table = exp.markdown_table
    mode_columns = [
        "mode",
        "tp",
        "fp",
        "fn",
        "precision",
        "recall",
        "f1",
        "alert_delay",
        "fps",
        "improvement_percent",
    ]
    env = json.loads((Path(rows[0]["run_dir"]) / "effective_config.json").read_text())
    cv_env = json.loads((Path(rows[0]["output_dir"]) / "opencv_environment.json").read_text())
    analyses = (ROOT / "docs/implementation_analysis.md").read_text()
    requirements = []
    for line in analyses.splitlines():
        if line.startswith("| ") and not line.startswith("| 기능"):
            cells = [s.strip() for s in line.strip("|").split("|")]
            if len(cells) == 7:
                name, status, file, fn, behavior, why, test = cells
                if name == "해상도별 FPS":
                    status, test = "IMPLEMENTED", "resolution_comparison.csv; 3×3 실제 실험"
                if name == "cv2.setNumThreads":
                    status, test = (
                        "EXPERIMENTAL / CONFIGURED",
                        "scripts/experiment_worker.py run; subprocess OFF/ON 실측",
                    )
                requirements.append(
                    dict(기능=name, 계획="최초 요구사항", 최종구현=status, 검증=test)
                )
    features = [
        (
            "Polygon",
            "영상마다 위험구역이 다르다",
            "경계를 재현 가능하게 고정한다",
            "원본 좌표 JSON과 schema·자기교차 검증; 마우스 편집",
            "src/zones.py",
            "edit_zones / save_zones / load_zones",
            "test_editor_keys_without_real_gui / test_json_and_display_coordinates",
        ),
        (
            "MOG2",
            "움직임 영역을 분리해야 한다",
            "객체 detector 없이 배경 대비 변화를 얻는다",
            "BGR MOG2 → threshold 200 → ellipse opening/closing → contour 면적 필터",
            "src/detect.py",
            "create_background_subtractor / detect_motion",
            "test_real_three_modes_csv_and_recording",
        ),
        (
            "Zone 판정",
            "움직임이 구역 밖에서도 발생한다",
            "구역 안 경보만 발생시킨다",
            "복원된 박스 하단 중앙점의 pointPolygonTest ≥ 0",
            "src/detect.py; src/preprocess.py",
            "check_intrusion / restore_boxes",
            "test_common_a_adapter_roi_coordinates",
        ),
        (
            "N-frame·State Machine",
            "한 프레임 노이즈가 경보를 만든다",
            "연속 관측으로 경보/해제를 확정한다",
            "기본 True 5회 ALERT; False 10회 CLEARED; 다음 관측에서 재진입 가능",
            "src/state.py; src/pipeline.py",
            "update_state / RealBackend.update",
            "test_reentry_immediately_after_cleared",
        ),
        (
            "Alert UI",
            "상태를 즉시 알아야 한다",
            "경보 가시성을 높인다",
            "원본 복사본에 polygon/box/상태·빨간 테두리·WARNING 표시",
            "src/pipeline.py",
            "_draw",
            "test_backend_default_threshold_and_warning",
        ),
        (
            "Event Recorder·Ring Buffer",
            "경보 후 녹화 시작만으로 이전 상황이 없다",
            "사건 전후 원본 증거를 보존한다",
            "시간 기준 deque·원본 복사·사건별 writer·CFR previous-frame hold",
            "src/recorder.py",
            "create_recorder / buffer_frame / record_frame / close_recorder",
            "test_three_second_buffer / test_event_file_and_long_extension",
        ),
        (
            "Pre/Post recording",
            "영상 시작/EOF에서 완전한 구간을 확보하지 못한다",
            "부족분을 숨기지 않는다",
            "alert 이전 3초, clear 이후 5초; 부족분 pre/post_truncated 표시; 없는 프레임 생성 금지",
            "src/recorder.py",
            "record_frame / _finish",
            "test_real_three_modes_csv_and_recording / test_irregular_hold_and_invalid_time",
        ),
        (
            "CSV logging",
            "여러 사건의 대응이 필요하다",
            "영상·구역·경보·클립을 식별한다",
            "EventLedger의 ID와 시각, duration, open_at_eof, video_path 기록",
            "src/pipeline.py",
            "EventLedger.observe / transitions / attach_clips / finish",
            "test_ledger_concurrent_dedup / test_failure_finalizes_open_writer",
        ),
    ]
    feature_text = "\n\n".join(
        f"### {name}\n\n문제: {problem}. 필요한 이유: {why}. 구현: {implementation}.\n\n관련 파일: `{file}`. 함수/클래스: `{fn}`. 검증: `{test}`."
        for name, problem, why, implementation, file, fn, test in features
    )
    failures = data["failures"]
    native = []
    for r in rows:
        if r["mode"] == "baseline" and r["repeat"] == 1:
            environment = json.loads((Path(r["run_dir"]) / "effective_config.json").read_text())
            native.append(
                dict(
                    video=r["video"],
                    native_size=str(environment["native_size"]),
                    source_fps=r["source_fps"],
                    frame_count=r["frame_count"],
                    eof_seconds=r["video_eof_time"],
                )
            )
    latency = numeric_rows(TAB / "performance_totals.csv")
    browse = [r for r in rows if r["video"] == "Browse.mp4" and r["repeat"] == 1]
    top = sorted(
        [
            r
            for r in data["stages"]
            if r["mode"] == "baseline" and r["stage"] not in ("total", "queue_wait", "display")
        ],
        key=lambda r: r["mean_ms"],
        reverse=True,
    )[0]
    sections = [
        "# OpenCV 위험구역 침입 감지 및 실시간 처리 최적화 — 최종 프로젝트 보고서",
        "## 1. 프로젝트 개요\n\n수동 CCTV 관찰 부담을 줄이기 위해 사용자가 지정한 위험구역에서 움직임 기반 침입을 감지하고 경보·사건 영상을 저장한다. 목표는 OpenCV 감지부터 사건 증거·정량 평가까지 연결하고, 실제 입력에서 처리 구조와 최적화의 효과 및 손실을 검증하는 것이다. 객체 종류/신원 식별 시스템은 아니다.",
        "## 2. 요구사항 대비 구현 결과\n\n상태는 실제 코드 기준이다. `EXPERIMENTAL / CONFIGURED`는 정식 기본 경로와 구별한다. 세부 파일·클래스·필요성은 [구현 분석](implementation_analysis.md)에 있다.\n\n"
        + table(requirements, ["기능", "계획", "최종구현", "검증"]),
        "## 3. 전체 시스템 Architecture\n\n```mermaid\nflowchart LR\n A[Video Input / FrameSource] --> B[Resize + Blur]\n B --> C[MOG2 + Morphology + Contour]\n C --> D[Coordinate restoration + Polygon]\n D --> E[N-frame / update_state]\n E --> F[Alert UI]\n E --> G[EventLedger CSV]\n A --> H[Original frame Ring Buffer]\n E --> I[Event Recorder]\n H --> I\n```\n\n영상 시각은 파일 CFR frame_index/source_fps이며 벽시계 처리시간과 분리한다. 원본 녹화와 분석 영상 좌표계도 분리한다. 영상 이탈은 intrusion 종료 조건이며 별도 분류 class가 아니다.",
        "## 4. 핵심 코드 구현\n\n" + feature_text,
        "## 5. 개발 과정에서 추가된 기능\n\nCLI(`src/main.py:main`), 설정 검증(`validate_config`), 실패 시 자원 해제, 실행별 UUID 디렉터리·유효 config·source SHA256, 단계 CSV, 실험 자동화, GT loader, 최대 cardinality 1:1 matching, Alert Delay, 그래프, regression tests가 연결되어 있다.\n\n이번 변경: `scripts/run_experiments.py:evaluation_intervals`는 원본 dict를 복사한 뒤 빈 종료 시각의 open_at_eof에만 `frame_count/source_fps`를 임시 종료로 적용한다. 원본 CSV/status/end_time은 보존한다. 내부 matcher 입력만 closed로 변환하며 매칭 내역은 원래 status와 evaluation_end_time을 함께 기록한다. start/alert가 비정상, 역전되거나 EOF가 없으면 review/N/A다. 정상 closed 사건의 기존 지연 누락 처리도 유지한다. `reevaluate_experiments.py`는 이전 실행을 새 평가 폴더로 재평가한다. `final_studies.py`는 3×3 및 5개 비교를 자동화한다. `experiment_worker.py`는 별도 subprocess에서 OpenCV 스레드 수를 설정한다. 감지·상태·녹화 AST는 포맷 전후 동일하다.\n\nBrowse 재평가 확인(최종 batch 첫 반복):\n\n"
        + table(
            browse,
            ["mode", "evaluation_status", "tp", "fp", "fn", "open_at_eof_count", "video_eof_time"],
        ),
        "## 6. 처리 구조\n\nBaseline: 동일 스레드에서 read→분석→상태→녹화. Threaded: 입력 producer와 분석 consumer가 bounded Queue로 연결되며 분석/녹화는 소비 경로에 남는다. Optimized: threaded 구조에 ROI/analysis downscale 옵션을 추가한다. 기본 ROI는 off이며 원본이 분석 상한보다 작으면 축소가 없으므로 optimized라는 이름 자체가 속도 개선을 보장하지 않는다. 세 모드 모두 동일 MOG2/상태 구현을 사용한다.",
        "## 7. 적용한 최적화 기법\n\n입력 스레드는 decode와 분석을 중첩한다. Downscale은 픽셀 수를 줄이고 좌표·면적을 보정한다. ROI는 다각형 전체 bounding rectangle+32px margin만 분석한다. Bounded Queue는 메모리/입력 선행량을 제한하며 실험은 용량 8→2 비교다(무제한 queue OFF가 아님). OpenCV threads는 라이브러리 내부 과도한 병렬화의 비용을 확인하기 위해 기본값→1로 비교했다. NumPy ROI min/max와 통계는 이미 벡터화되어 있으나 독립 OFF/ON 효과는 측정하지 않았다. 기대 효과를 실제 측정 효과로 취급하지 않는다.\n\n"
        + table(
            data["optimization"],
            [
                "technique",
                "off_fps",
                "on_fps",
                "improvement_percent",
                "off_f1",
                "on_f1",
                "delta_f1",
            ],
        ),
        "## 8. 실험 환경·재현 방법\n\n"
        + f"최종 batch: `{batch.relative_to(ROOT)}`. 비교 batch: `{study.relative_to(ROOT)}`.\n\nPython {env['python']}, OpenCV {env['opencv']}, NumPy {env['numpy']}; platform `{env['platform']}`; CPU `{env['processor'] or 'platform.processor 미제공'}`. OpenCV 기본 threads={cv_env['opencv_threads']}, OpenCL available={cv_env['opencl_available']}. GPU/OpenCL 성능은 N/A다.\n\n"
        + table(native, ["video", "native_size", "source_fps", "frame_count", "eof_seconds"])
        + "\n\n전체 7영상×3모드×3반복=63실행. 각 반복 모드 순서를 회전하고 EOF까지 처리한다. 녹화 on, no-display, queue block, interval=1, warmup=30 frames. 비교 실험은 Browse/Danger×2회이며 두 번째 반복의 조건 순서를 뒤집었다. GUI 비용·카메라 실시간 과부하·통계적 유의성 실험은 아니다.\n\n재현 명령:\n\n```bash\n.venv/bin/python -m pytest -q\n.venv/bin/ruff check .\n.venv/bin/ruff format --check .\n.venv/bin/python -m scripts.run_experiments --repeats 3\n.venv/bin/python -m scripts.final_studies --repeats 2\n.venv/bin/python -m scripts.build_final_report --batch <final-batch> --study <study-batch>\n```",
        "## 9. Ground Truth 기준\n\n확정 `docs/ground_truth_template.csv`를 수정하지 않는다. 실제 polygon 내부 점유 구간을 intrusion으로 정의하며 이탈은 종료 조건이다. basename과 JSON 파일명→단일 zone 이름만 명시적으로 대응한다. 침입 8구간/7영상이며 Near_walk는 전체 정상 marker다. 원본 GT snapshot과 audit CSV를 보존한다. 구간 overlap(끝점 접촉 포함), tolerance=0, 동일 영상·zone에서 기존 최대 1:1 matching을 사용한다. Open event의 평가용 구간은 [start_time, source EOF]다. 미매칭 경보는 FP, 미매칭 GT는 FN이다. TN이 정의되지 않아 일반 FPR은 N/A다.",
        "## 10. 감지 성능\n\n다음 전체 값은 각 모드 7영상×3반복의 micro 합산이다. 반복된 GT를 합산하므로 고유 사건 수와 다르며 모드끼리 합쳐 단일 정확도로 제시하지 않는다. Precision=TP/(TP+FP), Recall=TP/(TP+FN), F1=2TP/(2TP+FP+FN). Alert Delay는 TP의 alert_time−GT start 평균(초); 음수 보존, 시스템 latency와 구별한다. 분모가 0인 지표는 N/A다.\n\n"
        + table(modes, mode_columns)
        + "\n\n영상별 3반복 합산:\n\n"
        + table(
            videos,
            [
                "video",
                "mode",
                "tp",
                "fp",
                "fn",
                "precision",
                "recall",
                "f1",
                "alert_delay",
                "open_at_eof_count",
            ],
        ),
        "## 11. 오경보·미탐 분석\n\n아래는 baseline 첫 반복의 실제 미매칭 사건이다. `event_matches.csv`에서 모든 모드/반복을 추적할 수 있다. 구간만으로 객체의 신원이나 원인을 단정하지 않는다. 초기 MOG2 적응, 그림자/배경 변화, 정지 객체의 배경 흡수, 하단 중앙점과 실제 점유의 차이가 검토할 원인 후보다. EOF 미종료 사건도 예외 없이 포함되어 정답 밖 추가 경보가 FP로 반영된다.\n\n"
        + table(
            failures,
            [
                "video",
                "outcome",
                "event_id",
                "detected_start",
                "detected_end",
                "evaluation_end_time",
                "truth_start",
                "truth_end",
                "analysis",
            ],
        )
        + "\n\n![실제 backend 재실행 경보 화면](../outputs/final_report/figures/actual_alert.png)\n\n동결된 effective config로 RealBackend를 재실행한 화면이며 프레임/시각은 screenshot_provenance.json에 기록한다. 화면 FPS=0은 캡처용 표시값으로 성능 측정치가 아니다.",
        "## 12. Processing Performance\n\nFPS는 warmup 제외 analyzed_frames/elapsed_seconds이다. source FPS는 영상 시간축 정보다. 모드 FPS는 같은 21실행 처리량의 산술평균이며 pooled_fps는 합산 frames/합산 seconds다. 개선율=(new−baseline)/baseline×100. 초기화·해시·종료 시 CSV 쓰기/writer release는 측정 구간 밖이다.\n\n"
        + table(modes, ["mode", "fps", "pooled_fps", "fps_std", "improvement_percent"])
        + "\n\n"
        + table(
            latency,
            [
                "mode",
                "processed_frames",
                "measured_frames",
                "elapsed_seconds",
                "queue_wait_ms",
                "total_frame_ms",
                "dropped_frames",
            ],
        )
        + "\n\n![FPS](../outputs/final_report/figures/fps_comparison.png)",
        "## 13. 해상도별 성능\n\n원본 파일을 복제하지 않고 기존 input_resolution으로 입력 프레임을 192×144, 288×216, 384×288로 변환했다. baseline/threaded는 analysis_resolution 축소를 의도적으로 끄므로, 세 모드에서 동등하게 실제 크기를 바꾸기 위해 입력 크기 설정을 사용했다. 원본 녹화 크기는 유지한다. 구역 좌표도 기존 RealBackend가 변환한다. 이 입력 크기 실험은 min_area=200을 입력 픽셀 기준으로 유지하므로 크기 변화가 면적 필터에도 영향을 준다. 별도의 analysis downscale 기법 실험은 기존 scale 보정이 적용된다. Browse/Danger×2반복의 평균이며 고해상도 원본 일반화는 불가하다.\n\n"
        + table(data["resolution"], ["resolution", "mode", "runs", "fps", "tp", "fp", "fn", "f1"])
        + "\n\n![3×3](../outputs/final_report/figures/resolution_matrix.png)",
        "## 14. 최적화별 성능\n\n다섯 기법의 OFF/ON 비교는 각각 독립 새 실행이다. 입력 스레드=baseline→threaded; downscale=resize off→192×144; ROI=off→on; queue_capacity=8→2(block 유지); opencv_threads=환경 기본→1. 모두 원본 384×288, Browse/Danger×2반복, 동일 GT/녹화. 한 번에 하나의 옵션만 바꾸지만 환경 잡음과 반복 수 제약이 남는다. 음수 개선율도 그대로 기록한다.\n\n"
        + table(
            data["optimization"],
            [
                "technique",
                "off_fps",
                "on_fps",
                "improvement_percent",
                "off_tp",
                "off_fp",
                "off_fn",
                "on_tp",
                "on_fp",
                "on_fn",
                "delta_f1",
            ],
        )
        + "\n\n![최적화](../outputs/final_report/figures/optimization_effects.png)",
        "## 15. 병목 분석\n\n"
        + f"Baseline의 직접 처리 단계 중 가장 큰 평균은 `{top['stage']}` ({top['mean_ms']:.4f} ms)다. 다음 표는 실행별 stage 평균의 산술평균이다. queue_wait는 대기/백프레셔이며 CPU 연산이 아니다. total과 개별 stage를 중복 합산하지 않는다. threaded에서 read와 consumer가 중첩되므로 stage 합을 FPS 역수로 볼 수 없다. detect 내부 MOG2/morphology/contour별 시간은 분리 계측하지 않았다.\n\n"
        + table(data["stages"], ["mode", "stage", "mean_ms"])
        + "\n\n![병목](../outputs/final_report/figures/bottleneck.png)",
        "## 16. 시스템 한계\n\n- 움직임 기반이므로 사람/물체 식별·tracking이 없으며 장기 정지 침입이 배경에 흡수될 수 있다.\n- GT의 실제 점유와 박스 하단 중앙점 규칙은 다르다. 초기 경보와 중복 경보가 발생하며 overlap matching이 지연 품질을 보장하지 않는다.\n- 원본 보존 경로 없이 drop_oldest를 허용하지 않는다. bounded queue는 메모리를 제한하지만 실제 카메라 backend의 누적 지연을 제거한다고 보장할 수 없다. 실제 frame skipping도 비활성이다.\n- CLAHE/Gamma/Night Mode/Optical Flow/Heatmap/주의·금지 단계/화면 병목 표시는 미구현이다. OpenCL은 환경 미지원이다.\n- 저해상도 영상 7개, 비교 영상 2개이며 실시간 카메라·조명 변화·혼잡 상황을 대표하지 않는다. GUI 실제 마우스 조작은 자동화 대역 테스트이며 이번 최종 실험은 headless다.\n- pre는 최초 침입이 아닌 alert 기준이다. EOF에 남은 post 영상은 없으므로 truncated 표시는 정상이다. 원본 사건 종료값을 평가 편의상 쓰지 않는다.",
        "## 17. 향후 개선\n\nObject Detection과 Tracking으로 정지 객체·중복 경보 문제를 검증하고, adaptive threshold 및 저조도 전처리는 별도 영상/GT로 평가한다. 원본 녹화와 최신 프레임 분석을 분리한 후에만 drop_oldest·프레임 생략을 도입하고 원본 연속성 의미를 명확히 정의한다. 고해상도 원본, 실시간 입력 지연, 반복 수 확대 및 신뢰구간, GUI 비용을 추가 측정한다. 이는 현재 구현된 기능이 아니다.",
        "## 18. 최종 결론·교차검증\n\n"
        + table(modes, mode_columns)
        + f"\n\n핵심 위험구역 감지→경보→사건 녹화가 연결되어 있고 EOF 경보도 정상 평가된다. 속도 효과와 정확도 변화는 위 실측 그대로이며 감지 품질의 한계를 숨기지 않는다. 전체 {test_count()} tests 통과, repository-local Ruff 기본 correctness 검사/format 통과. 상위 사용자 Ruff 설정은 프로젝트 기준이 아니므로 ruff.toml로 기준을 고정했다. 기존 함수 구현은 AST 동등성을 확인했으며 포맷 변경만 있다.\n\n최종 GT와 기존 events.csv SHA256 보존을 검증했다. 보고서·PPT·그래프는 이 batch CSV에서 프로그램으로 생성하며 숫자를 수동 보정하지 않는다. CSV·그림의 SHA256은 frozen_results.json, 문서·PPT 검증은 validation.json에 기록한다. 재실행 성능은 달라질 수 있으며 이번 결과를 소급 덮어쓰지 않는다.",
    ]
    text = "\n\n".join(sections) + "\n"
    path = ROOT / "docs/final_experiment_report.md"
    path.write_text(text)
    return text


def architecture_figure():
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch

    fig, ax = plt.subplots(figsize=(12, 3))
    fig.patch.set_facecolor("#101F35")
    ax.set_facecolor("#101F35")
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 3)
    ax.axis("off")
    labels = [
        "Video input",
        "Resize / Blur",
        "MOG2\nMorphology",
        "Polygon\nN-frame",
        "State\nAlert",
        "CSV\nEvent clips",
    ]
    for i, label in enumerate(labels):
        x = 0.12 + i * 2
        ax.add_patch(
            FancyBboxPatch(
                (x, 0.9),
                1.65,
                1.2,
                boxstyle="round,pad=0.06",
                facecolor="#183750",
                edgecolor="#20A69A",
                linewidth=1.5,
            )
        )
        ax.text(x + 0.825, 1.5, label, ha="center", va="center", color="white", fontsize=12)
        if i < 5:
            ax.annotate(
                "",
                xy=(x + 1.97, 1.5),
                xytext=(x + 1.68, 1.5),
                arrowprops=dict(arrowstyle="->", color="#F0AB4B", lw=2),
            )
    ax.text(
        6,
        0.3,
        "Original frames → 3-second ring buffer → independent event recording → 5-second post-roll",
        ha="center",
        color="#AFC1D4",
        fontsize=11,
    )
    fig.tight_layout()
    fig.savefig(FIG / "architecture.png", dpi=180, facecolor=fig.get_facecolor())
    plt.close(fig)


def make_ppt(data):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    navy = "101F35"
    white = "F2F6FA"
    muted = "AFC1D4"
    teal = "20A69A"

    def box(slide, x, y, w, h, color):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        return shape

    def text(slide, words, x, y, w, h, size=22, color=white, bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.word_wrap = True
        for i, line in enumerate(words.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "NanumGothic"
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_after = Pt(12)
        return shape

    def slide(title, eyebrow="OPENCV / DANGER ZONE"):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(navy)
        box(s, 0.45, 0.45, 0.08, 0.65, teal)
        text(s, eyebrow, 0.7, 0.27, 11, 0.35, 11, teal, True)
        text(s, title, 0.7, 0.78, 12, 1, 30, white, True)
        text(s, "실제 코드 · 확정 GT · 실측 결과", 0.7, 7.02, 10, 0.3, 10, muted)
        text(s, f"{len(prs.slides):02d}", 12, 7.0, 0.6, 0.3, 11, muted)
        s.notes_slide.notes_text_frame.text = (
            "Source: "
            + str(data["batch"] / "summary")
            + "\nStudy: "
            + str(data["study"])
            + "\nDetails: docs/final_experiment_report.md; all figures/metrics generated from measured CSV."
        )
        return s

    def image(s, name, x, y, w):
        s.shapes.add_picture(str(FIG / name), Inches(x), Inches(y), width=Inches(w))

    def bullets(s, lines, x=0.8, y=2, w=11.6, size=23):
        text(s, "\n".join("• " + line for line in lines), x, y, w, min(4.6, 6.8 - y), size)

    def table(s, headers, records, x=0.8, y=2, w=11.7, h=3.0, size=16):
        t = s.shapes.add_table(
            len(records) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h)
        ).table
        for i, row in enumerate([headers] + records):
            for j, value in enumerate(row):
                cell = t.cell(i, j)
                cell.text = str(value)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string("236779" if i == 0 else "183047")
                for p in cell.text_frame.paragraphs:
                    p.font.name = "NanumGothic"
                    p.font.size = Pt(size)
                    p.font.bold = i == 0
                    p.font.color.rgb = RGBColor.from_string(white)

    ordered = [next(m for m in data["modes"] if m["mode"] == mode) for mode in exp.MODES]
    browse_first = next(
        r
        for r in data["rows"]
        if r["video"] == "Browse.mp4" and r["mode"] == "baseline" and r["repeat"] == 1
    )
    first_pass = aggregate(
        [r for r in data["rows"] if r["mode"] == "baseline" and r["repeat"] == 1]
    )
    s = slide("OpenCV 위험구역 침입 감지", "FINAL TEAM PROJECT")
    text(s, "실시간 처리 최적화", 0.8, 2.0, 11.5, 0.8, 32, teal, True)
    text(s, "OpenCV 기반 감지 → 경보 → 사건 기록 → 정량 검증", 0.8, 3.1, 11.8, 1, 25, muted)
    box(s, 0.8, 4.65, 11.7, 1.1, "183750")
    text(s, "7개 영상  ·  3개 모드  ·  3회 반복  ·  최종 63회 실행", 1.05, 4.88, 11, 0.7, 24)
    text(s, data["batch"].name, 0.8, 6.15, 12, 0.5, 13, muted)
    s = slide("01  문제 정의")
    bullets(
        s,
        [
            "CCTV의 위험구역을 계속 사람이 지켜보기 어렵다",
            "순간 움직임과 지속 침입을 구별해야 한다",
            "사건 전후 영상과 경보 이력을 함께 남겨야 한다",
            "빠른 처리와 감지 품질을 같은 실험에서 확인한다",
        ],
    )
    s = slide("02  프로젝트 목표와 범위")
    bullets(
        s,
        [
            "Polygon으로 사용자가 위험구역 지정",
            "MOG2 + N-frame + 상태 전이로 경보 확정",
            "경보 이전 3초 / 해제 이후 5초 원본 녹화",
            "Baseline · Threaded · Optimized 성능 비교",
            "객체 식별·Tracking·저조도 보정은 후속 과제",
        ],
    )
    s = slide("03  실제 시스템 Architecture")
    image(s, "architecture.png", 0.65, 2, 12)
    text(
        s,
        "입력 시간축과 처리 벽시계, 원본 녹화와 분석 좌표계를 분리",
        0.85,
        5.7,
        11.7,
        0.8,
        23,
        muted,
    )
    s = slide("04  Polygon + Motion Detection")
    image(s, "zone_view.png", 0.8, 2.05, 5.5)
    bullets(
        s,
        [
            "edit_zones / save_zones / load_zones",
            "MOG2 → 그림자 제외 → Morphology",
            "Contour 면적 필터 → 하단 중앙점",
            "pointPolygonTest로 내부·경계 판정",
        ],
        6.7,
        2.15,
        5.8,
        20,
    )
    s = slide("05  N-frame + State Machine")
    text(s, "IDLE → DETECTING → ALERT → CLEARED", 0.8, 2, 11.6, 1, 29, teal, True)
    bullets(
        s,
        [
            "기본 5프레임 연속 침입 → ALERT",
            "기본 10프레임 연속 비침입 → CLEARED",
            "구역별 상태와 사건 ID를 독립 관리",
            "update_state / RealBackend.update",
        ],
        y=3.25,
        size=23,
    )
    s = slide("06  Alert + Event Recording")
    image(s, "actual_alert.png", 0.8, 2.05, 5.5)
    bullets(
        s,
        [
            "3초 pre → 경보 구간 → 5초 post",
            "deque 원본 버퍼 + 사건별 writer",
            "CSV의 사건·구역·시각·클립 연결",
            "EOF 부족 구간은 truncated 표시",
        ],
        6.7,
        2.15,
        5.8,
        20,
    )
    text(s, "실제 backend 재실행 화면 / FPS=0은 캡처 표시값", 0.8, 6.4, 11.5, 0.4, 12, muted)
    s = slide("07  구현 결과와 추가 자동화")
    table(
        s,
        ["완성된 핵심", "검증·자동화"],
        [
            ["Polygon / MOG2 / Zone", "CLI + config validation"],
            ["N-frame / State / Alert", "GT loader + 1:1 matching"],
            ["3초/5초 녹화 + CSV", "EOF 평가 adapter + 회귀 테스트"],
            ["세 모드 + bounded Queue", "해상도 3×3 + 최적화 5종 비교"],
        ],
        h=3.6,
        size=19,
    )
    text(
        s,
        "drop_oldest·실제 frame skipping은 녹화/연속성 제약으로 비활성",
        0.8,
        6.15,
        11.7,
        0.7,
        19,
        muted,
    )
    s = slide("08  Baseline vs Threaded vs Optimized")
    table(
        s,
        ["모드", "실제 실행 구조"],
        [
            ["Baseline", "순차 입력 → 분석 → 상태 → 녹화"],
            ["Threaded", "입력 producer + bounded Queue + consumer"],
            ["Optimized", "Threaded 구조 + 선택적 ROI / 분석 축소"],
        ],
        h=3,
        size=20,
    )
    text(
        s,
        "동일 감지 코드 사용 / 기본 384×288에서는 축소가 적용되지 않음",
        0.8,
        5.6,
        11.6,
        1,
        21,
        muted,
    )
    s = slide("09  최적화 실측: 속도와 정확도 함께")
    table(
        s,
        ["기법", "OFF FPS", "ON FPS", "변화 %", "F1 변화"],
        [
            [
                r["technique"],
                fmt(r["off_fps"], 1),
                fmt(r["on_fps"], 1),
                fmt(r["improvement_percent"], 2),
                fmt(r["delta_f1"], 3),
            ]
            for r in data["optimization"]
        ],
        h=3.7,
        size=16,
    )
    text(
        s,
        "Browse/Danger × 2회 / queue는 8→2 / OpenCV threads는 기본→1",
        0.8,
        6.2,
        11.7,
        0.6,
        17,
        muted,
    )
    s = slide("10  실험 설계 / 확정 Ground Truth")
    bullets(
        s,
        [
            "전체: 7영상 × 3모드 × 3회 = 63실행",
            "확정 GT: 고유 침입 8구간 / GT 원본 변경 없음",
            "같은 영상·Zone의 overlap → 최대 1:1 매칭",
            "open_at_eof도 [start, frame_count/source_fps]로 평가",
            "원본 events.csv의 빈 end_time과 status는 보존",
        ],
        size=22,
    )
    s = slide("11  감지 성능: 과장 없이 공개")
    table(
        s,
        ["모드", "TP / FP / FN", "Precision", "Recall", "F1", "Delay(s)"],
        [
            [
                m["mode"],
                f"{int(m['tp'])} / {int(m['fp'])} / {int(m['fn'])}",
                fmt(m["precision"]),
                fmt(m["recall"]),
                fmt(m["f1"]),
                fmt(m["alert_delay"]),
            ]
            for m in ordered
        ],
        h=2.5,
        size=17,
    )
    text(
        s,
        f"각 모드 3반복 합산 / baseline 1회: TP {int(first_pass['tp'])} · FP {int(first_pass['fp'])} · FN {int(first_pass['fn'])}",
        0.8,
        5.2,
        11.7,
        0.7,
        20,
        muted,
    )
    text(
        s,
        f"Browse: TP {int(browse_first['tp'])} / FP {int(browse_first['fp'])} / FN {int(browse_first['fn'])} (각 모드·반복), EOF 경보 포함",
        0.8,
        6,
        11.7,
        0.7,
        19,
        teal,
    )
    s = slide("12  처리 성능 / 해상도 비교")
    table(
        s,
        ["모드", "Processing FPS", "Baseline 대비 %"],
        [[m["mode"], fmt(m["fps"], 2), fmt(m["improvement_percent"], 2)] for m in ordered],
        x=0.7,
        y=1.9,
        w=5.8,
        h=2.3,
        size=17,
    )
    image(s, "resolution_matrix.png", 6.8, 1.9, 5.9)
    text(
        s,
        "Source FPS와 처리 throughput은 다른 값\nWarmup 30 frame 제외 / 녹화 포함 / GUI 제외",
        0.8,
        4.8,
        11.7,
        1.25,
        21,
        muted,
    )
    s = slide("13  실패 사례와 한계")
    bullets(
        s,
        [
            "정답 밖 추가 경보와 GT 미매칭이 실제로 발생",
            "움직임 기반: 정지 객체·배경 변화에 취약할 수 있음",
            "하단 중앙점 기준과 실제 polygon 점유는 차이가 있음",
            "bounded Queue ≠ 카메라 지연 누적 완전 제거",
            "OpenCL 미지원 / 7개 저해상도 영상으로 일반화 제한",
        ],
        size=22,
    )
    s = slide("14  결론 / 다음 단계")
    bullets(
        s,
        [
            "핵심 감지·경보·원본 사건 녹화 경로 검증",
            "EOF 경보 평가 복구 + 재현 가능한 결과 고정",
            f"Baseline {ordered[0]['fps']:.2f} FPS / F1 {ordered[0]['f1']:.3f}",
            f"{test_count()} tests + lint/format 통과; 수치 출처는 최종 CSV",
            "다음: 객체 검출·Tracking·적응형 임계값·고해상도 검증",
        ],
        size=22,
    )
    path = OUT / "OpenCV_Danger_Zone_Final_Presentation.pptx"
    prs.save(path)
    return path


def make_pdf(markdown):
    """Render the full report with Korean fonts and wrapped tables; keep Markdown intact."""
    import re
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
        Image,
        Preformatted,
    )

    pdfmetrics.registerFont(TTFont("Nanum", "/usr/share/fonts/truetype/nanum/NanumGothic.ttf"))
    pdfmetrics.registerFont(
        TTFont("NanumBold", "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf")
    )
    styles = getSampleStyleSheet()
    for name in ("Normal", "Title", "Heading1", "Heading2", "Heading3"):
        styles[name].fontName = "NanumBold" if name != "Normal" else "Nanum"
        styles[name].wordWrap = "CJK"
    styles["Normal"].fontSize = 9
    styles["Normal"].leading = 14
    cellstyle = ParagraphStyle(
        "Cell", fontName="Nanum", fontSize=6.6, leading=9, wordWrap="CJK", alignment=TA_LEFT
    )
    code = ParagraphStyle("CodeSmall", fontName="Nanum", fontSize=7, leading=10)
    width = A4[0] - 72
    story = []
    lines = markdown.splitlines()
    i = 0

    def clean(s):
        s = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", s)
        return escape(s.replace("`", ""))

    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            language = line[3:]
            content = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                content.append(lines[i])
                i += 1
            if language == "mermaid":
                story.append(Image(str(FIG / "architecture.png"), width=width, height=width / 4))
            else:
                story.append(Preformatted("\n".join(textwrap.fill(t, 95) for t in content), code))
        elif line.startswith("| "):
            data = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s:|\-]+$", lines[i]):
                    data.append(
                        [
                            Paragraph(clean(c.strip()), cellstyle)
                            for c in lines[i].strip("|").split("|")
                        ]
                    )
                i += 1
            table = Table(
                data, colWidths=[width / len(data[0])] * len(data[0]), repeatRows=1, hAlign="LEFT"
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCEEF1")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#BAC7CC")),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.extend([table, Spacer(1, 10)])
            continue
        elif line.startswith("!["):
            match = re.search(r"\]\(([^)]+)\)", line)
            path = (ROOT / "docs" / match.group(1)).resolve()
            from PIL import Image as PILImage

            with PILImage.open(path) as img:
                w, h = img.size
            story.append(Image(str(path), width=width, height=width * h / w))
        elif line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            style = "Title" if level == 1 else "Heading1" if level == 2 else "Heading2"
            heading = clean(line.lstrip("# "))
            if level == 1:
                heading = heading.replace(" — ", "<br/>")
            story.append(Paragraph(heading, styles[style]))
        elif line.strip():
            story.append(Paragraph(clean(line), styles["Normal"]))
            story.append(Spacer(1, 5))
        i += 1
    path = OUT / "OpenCV_Danger_Zone_Final_Report.pdf"
    doc = SimpleDocTemplate(
        str(path), pagesize=A4, rightMargin=36, leftMargin=36, topMargin=40, bottomMargin=40
    )

    def footer(canvas, document):
        canvas.setFont("Nanum", 8)
        canvas.drawRightString(A4[0] - 36, 22, str(document.page))

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return path


def validate_artifacts(data, ppt, pdf, markdown):
    from pptx import Presentation

    prs = Presentation(ppt)
    assert len(prs.slides) == 15
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 10
            assert shape.top + shape.height <= prs.slide_height + 10
            if shape.has_text_frame:
                texts.append(shape.text)
            if shape.has_table:
                texts.extend(c.text for row in shape.table.rows for c in row.cells)
    combined = "\n".join(texts)
    for mode in data["modes"]:
        for key in ("precision", "recall", "f1", "alert_delay"):
            assert fmt(mode[key]) in combined
            assert fmt(mode[key], 4) in markdown
        assert fmt(mode["fps"], 2) in combined
    for row in data["optimization"]:
        assert fmt(row["off_fps"], 1) in combined and fmt(row["on_fps"], 1) in combined
    frozen = json.loads((TAB / "frozen_results.json").read_text())
    assert all(
        hashlib.sha256(Path(p).read_bytes()).hexdigest() == digest for p, digest in frozen.items()
    )
    validation = dict(
        main_runs=len(data["rows"]),
        study_runs=76,
        all_completed=True,
        all_evaluated=True,
        source_frames_complete=True,
        recording_verified=True,
        input_hashes_preserved=True,
        frozen_results_unchanged=True,
        slides=len(prs.slides),
        slide_bounds_checked=True,
        csv_markdown_ppt_metrics_checked=True,
        pdf_bytes=pdf.stat().st_size,
        note="PPT XML/text/bounds checked; no PowerPoint/LibreOffice renderer available for pixel rendering",
    )
    (TAB / "validation.json").write_text(json.dumps(validation, indent=2))
    return validation


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--batch", type=Path, required=True)
    parser.add_argument("--study", type=Path, required=True)
    args = parser.parse_args(argv)
    OUT.mkdir(parents=True, exist_ok=True)
    FIG.mkdir(exist_ok=True)
    TAB.mkdir(exist_ok=True)
    data = prepare(args.batch.resolve(), args.study.resolve())
    architecture_figure()
    # Freeze inputs/derived numeric tables before report and slides are generated.
    paths = set()
    for root in (data["batch"], data["study"], TAB, FIG):
        paths.update(
            p
            for p in root.rglob("*")
            if p.suffix in (".csv", ".json", ".png")
            and p.name not in ("frozen_results.json", "validation.json")
        )
    hashes = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(paths)}
    (TAB / "frozen_results.json").write_text(json.dumps(hashes, indent=2))
    markdown = make_markdown(data)
    ppt = make_ppt(data)
    pdf = make_pdf(markdown)
    print(json.dumps(validate_artifacts(data, ppt, pdf, markdown), indent=2))


if __name__ == "__main__":
    raise SystemExit(main())
